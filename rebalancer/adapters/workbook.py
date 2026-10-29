"""Excel workbook adapter: the only place that knows the 7-sheet XLSX format (doc 07 §3, doc 08 §12).

Returns a validated ``Dataset`` plus a list of issues with sheet/cell references. Global errors make the
whole file unusable; per-client errors block only that client.
"""

import io
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

from dataclasses import replace

from ..engine.model import CASH, CZK, FO, SRO, CashBalance, Client, Dataset, Flow, Instrument, Lot, Policy
from ..engine.policy import default_policy, validate_policy
from ..engine.results import Issue

MAX_UPLOAD = 2 * 1024 * 1024
MAX_UNCOMPRESSED = 20 * 1024 * 1024
MAX_CELLS = 10_000
MAX_LOTS = 1_000
MAX_FLOWS = 500
MAX_UNITS = 10**9
SHEETS = ["Instrumenty", "Kurzy", "Klienti", "Pozice", "Hotovost", "Cash flow"]
CZ_MONTH_DATE = re.compile(r"(\d{1,2})\.\s*(\d{1,2})\.\s*(\d{4})")


class WorkbookError(Exception):
    """The file cannot be read at all (global error)."""

    def __init__(self, issues: list[Issue]):
        super().__init__("; ".join(i.message for i in issues))
        self.issues = issues


def _num(v) -> Decimal | None:
    """Excel doubles normalised to ≤ 6 decimal places (DI-1: 78.40000000000001 → 78.4)."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        d = Decimal(repr(v)) if isinstance(v, float) else Decimal(v)
    elif isinstance(v, str):
        try:
            d = Decimal(v.strip().replace(" ", "").replace(" ", "").replace(",", "."))
        except InvalidOperation:
            return None
    else:
        return None
    if not d.is_finite() or abs(d) >= Decimal(10) ** 13:
        return None
    try:
        d = d.quantize(Decimal("0.000001")).normalize()
    except InvalidOperation:
        return None
    if d == 0:
        return Decimal(0)
    # keep plain notation on the wire: normalize() turns 50 into 5E+1
    return d.quantize(Decimal(1)) if d == d.to_integral_value() else d


def _date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        m = CZ_MONTH_DATE.search(v)
        if m:
            try:
                return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
            except ValueError:
                return None
        try:
            return date.fromisoformat(v.strip())
        except ValueError:
            return None
    return None


def _text(v) -> str:
    return "" if v is None else str(v).strip()


# Optional sheet "Strategie": one row per client; target weights stay on "Klienti".
STRATEGY_COLUMNS: list[tuple[str, str, bool]] = [  # (header, Policy field, integer) — the editable parameters
    ("Absolutní tolerance (p. b.)", "band_abs", False),
    ("Relativní tolerance (% cíle)", "band_rel", False),
    ("Minimální obchod (Kč)", "min_trade", False),
    ("Max. doba čekání (dní)", "max_wait_days", True),
]
SETTLEMENT_HEADER = "Vypořádání (pracovní dny)"  # optional column on "Instrumenty"; default T+2


def read_workbook(data: bytes, source: str = "upload.xlsx") -> tuple[Dataset, list[Issue]]:
    ds, issues, _ = read_book(data, source)
    return ds, issues


def read_book(data: bytes, source: str = "upload.xlsx") -> tuple[Dataset, list[Issue], dict[str, Policy]]:
    """Dataset + issues + client strategies (from the optional "Strategie" sheet, else the defaults)."""
    issues: list[Issue] = []

    def err(code, msg, sheet=None, cell=None, client=None, severity="error"):
        issues.append(Issue(severity, code, msg, sheet, cell, client))

    if len(data) > MAX_UPLOAD:
        raise WorkbookError([Issue("error", "too_large", "Soubor je větší než 2 MiB.")])
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_UNCOMPRESSED:
                raise WorkbookError([Issue("error", "too_large", "Rozbalený obsah souboru je větší než 20 MiB.")])
        wb = load_workbook(io.BytesIO(data), data_only=True)
        wbf = load_workbook(io.BytesIO(data), data_only=False)
    except WorkbookError:
        raise
    except Exception:
        raise WorkbookError([Issue("error", "not_xlsx", "Soubor není platný sešit Excel (.xlsx).")])

    missing = [s for s in SHEETS if s not in wb.sheetnames]
    if missing:
        raise WorkbookError([Issue("error", "missing_sheet", f"Chybí list „{s}“.", s) for s in missing])

    for ws in wb.worksheets:
        if ws.max_row * ws.max_column > MAX_CELLS:
            raise WorkbookError([Issue("error", "too_many_cells", f"List „{ws.title}“ má více než {MAX_CELLS} buněk.",
                                       ws.title)])

    def rows(sheet):
        ws = wb[sheet]
        for r in ws.iter_rows(min_row=2):
            if all(c.value is None for c in r):
                continue
            for c in r:
                if wbf[sheet][c.coordinate].data_type == "f":
                    err("formula", "Buňka obsahuje vzorec; použije se uložená vypočtená hodnota.",
                        sheet, c.coordinate, severity="warning")
            yield r

    # valuation date: „Čti mě“ text „Ocenění k datu: 23. 10. 2026“
    valuation = None
    if "Čti mě" in wb.sheetnames:
        for r in wb["Čti mě"].iter_rows():
            for c in r:
                if isinstance(c.value, str) and "Ocenění k datu" in c.value:
                    valuation = _date(c.value)
    if valuation is None:
        raise WorkbookError([Issue("error", "no_valuation_date",
                                   "Nenalezeno datum ocenění (list „Čti mě“, text „Ocenění k datu: D. M. RRRR“).",
                                   "Čti mě")])

    instruments: list[Instrument] = []
    inst_head = [_text(c.value) for c in wb["Instrumenty"][1]]
    settle_col = inst_head.index(SETTLEMENT_HEADER) if SETTLEMENT_HEADER in inst_head else None
    for r in rows("Instrumenty"):
        t, name, isin, ccy, cls, price = (r[i].value if i < len(r) else None for i in range(6))
        p = _num(price)
        row = r[0].row
        if not _text(t):
            err("missing_value", "Chybí ticker.", "Instrumenty", f"A{row}")
            continue
        if p is None or p <= 0:
            err("bad_price", f"Neplatná cena instrumentu {t}.", "Instrumenty", f"F{row}")
            continue
        if _text(ccy) == "":
            err("missing_value", f"Chybí měna instrumentu {t}.", "Instrumenty", f"D{row}")
            continue
        if any(i.ticker == _text(t) for i in instruments):
            err("duplicate", f"Duplicitní ticker {t}.", "Instrumenty", f"A{row}")
            continue
        settle = 2
        if settle_col is not None and settle_col < len(r) and r[settle_col].value is not None:
            sv = _num(r[settle_col].value)
            if sv is None or sv != sv.to_integral_value() or not (0 <= sv <= 10):
                err("bad_settlement", f"Neplatná lhůta vypořádání instrumentu {t} (0–10 pracovních dnů).",
                    "Instrumenty", r[settle_col].coordinate)
                continue
            settle = int(sv)
        instruments.append(Instrument(_text(t), _text(name), _text(isin), _text(ccy).upper(), _text(cls), p, settle))
    if not instruments:
        err("no_instruments", "List „Instrumenty“ neobsahuje žádný platný instrument.", "Instrumenty")

    fx: dict[str, Decimal] = {CZK: Decimal(1)}
    for r in rows("Kurzy"):
        pair, rate = r[0].value, r[1].value if len(r) > 1 else None
        row = r[0].row
        m = re.fullmatch(r"\s*([A-Z]{3})\s*/\s*CZK\s*", _text(pair).upper())
        v = _num(rate)
        if not m:
            err("bad_fx_pair", f"Neplatný měnový pár „{_text(pair)}“ (očekáváno XXX/CZK).", "Kurzy", f"A{row}")
            continue
        if m.group(1) in fx:
            err("duplicate", f"Duplicitní kurz {m.group(1)}/CZK.", "Kurzy", f"A{row}")
            continue
        if v is None or v <= 0:
            err("bad_fx", f"Neplatný kurz {m.group(1)}.", "Kurzy", f"B{row}")
            continue
        fx[m.group(1)] = v
    for i in instruments:
        if i.currency not in fx:
            err("missing_fx", f"Chybí kurz pro měnu {i.currency} (instrument {i.ticker}).", "Kurzy")
    if any(i.severity == "error" for i in issues):
        raise WorkbookError([i for i in issues if i.severity == "error"])
    tickers = [i.ticker for i in instruments]

    header = [_text(c.value) for c in wb["Klienti"][1]]
    target_cols: dict[str, int] = {}
    for idx, h in enumerate(header):
        m = re.fullmatch(r"Cíl\s+(.+?)\s*%", h)
        if m:
            key = m.group(1).strip()
            target_cols[CASH if key.lower() == "hotovost" else key] = idx
    for t in tickers + [CASH]:
        if t not in target_cols:
            err("missing_column", f"Chybí sloupec „Cíl {t if t != CASH else 'hotovost'} %“.", "Klienti", "1")
    if any(i.severity == "error" for i in issues):
        raise WorkbookError([i for i in issues if i.severity == "error"])

    raw_clients = []
    names = set()
    for r in rows("Klienti"):
        row = r[0].row
        name = _text(r[0].value)
        if not name:
            err("missing_value", "Chybí jméno klienta.", "Klienti", f"A{row}")
            continue
        if name in names:
            err("duplicate", f"Duplicitní klient „{name}“.", "Klienti", f"A{row}")
            continue
        names.add(name)
        regime = _text(r[2].value)
        if regime not in (FO, SRO):
            err("bad_regime", f"Neznámý daňový režim „{regime}“ (FO nebo s.r.o.).", "Klienti", f"C{row}", name)
        targets = {}
        for k, idx in target_cols.items():
            v = _num(r[idx].value) if idx < len(r) else None
            col = r[idx].column_letter if idx < len(r) else "?"
            if v is None or v < 0 or v > 100:
                err("bad_target", f"Neplatná cílová váha {k if k != CASH else 'hotovost'}.", "Klienti", f"{col}{row}", name)
                v = Decimal(0)
            targets[k] = v
        if sum(targets.values(), Decimal(0)) != 100:
            err("targets_sum", f"Cílové váhy klienta {name} nedávají součet 100 % "
                f"({sum(targets.values(), Decimal(0))} %).", "Klienti", f"D{row}", name)
        note = _text(r[8].value) if len(r) > 8 else ""
        raw_clients.append(dict(name=name, profile=_text(r[1].value), regime=regime, targets=targets, note=note))

    lots: dict[str, list[Lot]] = {c["name"]: [] for c in raw_clients}
    n_lots = 0
    for r in rows("Pozice"):
        row = r[0].row
        name, t, d, n, p = (r[i].value if i < len(r) else None for i in range(5))
        name = _text(name)
        n_lots += 1
        if name not in lots:
            err("orphan", f"Pozice patří neznámému klientovi „{name}“.", "Pozice", f"A{row}")
            continue
        if _text(t) not in tickers:
            err("unknown_ticker", f"Neznámý ticker „{_text(t)}“.", "Pozice", f"B{row}", name)
            continue
        dd = _date(d)
        if dd is None or dd > valuation:
            err("bad_date", "Neplatné datum nákupu (chybí nebo je po datu ocenění).", "Pozice", f"C{row}", name)
            continue
        units = _num(n)
        if units is None or units <= 0 or units != units.to_integral_value() or units > MAX_UNITS:
            err("bad_units", "Počet kusů musí být kladné celé číslo.", "Pozice", f"D{row}", name)
            continue
        price = _num(p)
        if price is None or price <= 0:
            err("bad_price", "Neplatná nákupní cena.", "Pozice", f"E{row}", name)
            continue
        lot = Lot(_text(t), dd, int(units), price)
        if lot in lots[name]:
            err("duplicate", "Stejný nákup (datum, kusy, cena) je uveden dvakrát – zkontrolujte, zda nejde o zkopírovaný řádek.",
                "Pozice", f"A{row}", name, severity="warning")
        lots[name].append(lot)
    if n_lots > MAX_LOTS:
        raise WorkbookError([Issue("error", "too_many_lots", f"Více než {MAX_LOTS} pozic.", "Pozice")])

    cash: dict[str, dict[str, Decimal]] = {c["name"]: {} for c in raw_clients}
    for r in rows("Hotovost"):
        row = r[0].row
        name, ccy, amt = (_text(r[0].value), _text(r[1].value).upper(), _num(r[2].value) if len(r) > 2 else None)
        if name not in cash:
            err("orphan", f"Hotovost patří neznámému klientovi „{name}“.", "Hotovost", f"A{row}")
            continue
        if ccy not in fx:
            err("missing_fx", f"Měna {ccy} nemá kurz.", "Hotovost", f"B{row}", name)
            continue
        if amt is None:
            err("bad_amount", "Neplatný zůstatek.", "Hotovost", f"C{row}", name)
            continue
        if amt < 0:
            err("negative_cash", "Záporný zůstatek hotovosti.", "Hotovost", f"C{row}", name)
            continue
        if ccy in cash[name]:
            err("duplicate", f"Duplicitní zůstatek v měně {ccy}.", "Hotovost", f"B{row}", name)
            continue
        cash[name][ccy] = amt

    flows: dict[str, list[Flow]] = {c["name"]: [] for c in raw_clients}
    n_flows = 0
    for r in rows("Cash flow"):
        row = r[0].row
        name, d, kind, amt, note = (r[i].value if i < len(r) else None for i in range(5))
        name = _text(name)
        n_flows += 1
        if name not in flows:
            err("orphan", f"Cash flow patří neznámému klientovi „{name}“.", "Cash flow", f"A{row}")
            continue
        dd, a = _date(d), _num(amt)
        if dd is None:
            err("bad_date", "Neplatné datum cash flow.", "Cash flow", f"B{row}", name)
            continue
        if a is None or a == 0:
            err("bad_amount", "Neplatná částka cash flow.", "Cash flow", f"D{row}", name)
            continue
        k = _text(kind)
        if (k.lower() == "vklad" and a < 0) or (k.lower() == "výběr" and a > 0):
            err("sign_mismatch", f"Typ „{k}“ neodpovídá znaménku částky.", "Cash flow", f"D{row}", name)
            continue
        if dd <= valuation:
            err("past_flow", "Cash flow k datu ocenění nebo dříve se ignoruje (předpokládá se, že je v hotovosti).",
                "Cash flow", f"B{row}", name, severity="warning")
        flows[name].append(Flow(dd, a, k, _text(note)))
    if n_flows > MAX_FLOWS:
        raise WorkbookError([Issue("error", "too_many_flows", f"Více než {MAX_FLOWS} cash flow.", "Cash flow")])

    clients = []
    for c in raw_clients:
        n = c["name"]
        if not lots[n] and not cash[n]:
            err("empty_client", f"Klient {n} nemá žádné pozice ani hotovost.", "Klienti", None, n)
        clients.append(Client(n, c["profile"], c["regime"], c["targets"], c["note"], lots[n],
                              [CashBalance(k, v) for k, v in cash[n].items()], flows[n]))
    if not clients:
        raise WorkbookError([Issue("error", "no_clients", "Sešit neobsahuje žádného klienta.", "Klienti")])
    policies = {c.name: default_policy(c.targets) for c in clients}
    if "Strategie" in wb.sheetnames:
        head = [_text(c.value) for c in wb["Strategie"][1]]
        cols = {field: head.index(h) for h, field, _ in STRATEGY_COLUMNS if h in head}
        for r in rows("Strategie"):
            row = r[0].row
            name = _text(r[0].value)
            if name not in policies:
                err("orphan", f"Strategie patří neznámému klientovi „{name}“.", "Strategie", f"A{row}")
                continue
            values = {}
            for h, field, integer in STRATEGY_COLUMNS:
                idx = cols.get(field)
                if idx is None or idx >= len(r) or r[idx].value is None:
                    continue
                cell = r[idx].coordinate
                v = _num(r[idx].value)
                if v is None or (integer and v != v.to_integral_value()):
                    err("bad_policy", f"Neplatná hodnota „{h}“.", "Strategie", cell, name)
                    continue
                values[field] = int(v) if integer else v
            policies[name] = replace(policies[name], **values)
            for p in validate_policy(policies[name], tickers):
                col = next((h for h, f, _ in STRATEGY_COLUMNS if f == p.field), p.field)
                err("bad_policy", f"Strategie: „{col}“ mimo povolený rozsah.", "Strategie", f"A{row}", name)
    # an incomplete file (orphan rows, duplicate clients, …) is rejected as a whole (owner decision 2)
    global_errors = [i for i in issues if i.severity == "error" and i.client is None]
    if global_errors:
        raise WorkbookError(global_errors)
    return Dataset(valuation, instruments, fx, clients, source), issues, policies
