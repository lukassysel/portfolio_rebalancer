"""API contract and portability tests (doc 10 §1: CLI == API, decimal strings on the wire, import rejection)."""

import io
import json

from fastapi.testclient import TestClient
from datetime import datetime

from openpyxl import load_workbook

from rebalancer.adapters import demo_bytes
from rebalancer.api.app import app
from rebalancer.cli.main import evaluate_file

client = TestClient(app)


def test_health():
    r = client.get("/api/health")
    assert r.status_code == 200 and r.json()["clients"] == 5


def test_demo_snapshot_and_decimal_strings():
    r = client.get("/api/demo").json()
    assert r["demo"] is True and len(r["result"]["clients"]) == 5
    a = r["result"]["clients"][0]
    assert isinstance(a["valuation"]["nav"], str) and a["valuation"]["nav"].startswith("10267798.2")
    assert isinstance(r["dataset"]["instruments"][0]["price"], str)


def test_evaluate_roundtrip_matches_demo():
    snap = client.get("/api/demo").json()
    r = client.post("/api/evaluate", json={"dataset": snap["dataset"], "policies": snap["policies"],
                                           "issues": snap["issues"], "model_version": snap["model_version"]})
    assert r.status_code == 200
    assert r.json() == snap["result"]


def test_cli_equals_api():
    snap = client.get("/api/demo").json()
    assert json.loads(evaluate_file(None)) == snap["result"]


def test_policy_edit_recalculates():
    snap = client.get("/api/demo").json()
    snap["policies"]["Klient C"]["band_abs"] = "0.4"
    r = client.post("/api/evaluate", json={"dataset": snap["dataset"], "policies": snap["policies"]}).json()
    c = next(x for x in r["clients"] if x["name"] == "Klient C")
    assert c["recommended"].startswith("NOW_")


def test_model_version_mismatch():
    snap = client.get("/api/demo").json()
    r = client.post("/api/evaluate", json={"dataset": snap["dataset"], "policies": snap["policies"],
                                           "model_version": "0.0.0"})
    assert r.status_code == 409


def test_import_valid_modified_copy():
    wb = load_workbook(io.BytesIO(demo_bytes()))
    wb["Instrumenty"]["F2"] = 156.75  # VWCE +10 %
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post("/api/import", files={"file": ("mesic2.xlsx", buf.getvalue())})
    assert r.status_code == 200
    body = r.json()
    assert body["demo"] is False and body["dataset"]["instruments"][0]["price"] == "156.75"


def test_import_rejects_garbage():
    r = client.post("/api/import", files={"file": ("x.xlsx", b"not a workbook")})
    assert r.status_code == 422 and r.json()["issues"][0]["code"] == "not_xlsx"


def test_import_rejects_missing_sheet():
    wb = load_workbook(io.BytesIO(demo_bytes()))
    del wb["Kurzy"]
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post("/api/import", files={"file": ("x.xlsx", buf.getvalue())})
    assert r.status_code == 422 and r.json()["issues"][0]["sheet"] == "Kurzy"


def test_import_client_level_error_isolated():
    wb = load_workbook(io.BytesIO(demo_bytes()))
    wb["Hotovost"]["C6"] = -5  # Klient C negative cash
    buf = io.BytesIO()
    wb.save(buf)
    body = client.post("/api/import", files={"file": ("x.xlsx", buf.getvalue())}).json()
    res = {c["name"]: c for c in body["result"]["clients"]}
    assert res["Klient C"]["outcome"]["code"] == "fix_data"
    assert res["Klient C"]["issues"][0]["cell"] == "C6"
    assert res["Klient A"]["status"] == "ok"


def test_unknown_api_route_is_json_404():
    r = client.get("/api/nope")
    assert r.status_code == 404 and r.json()["detail"]


def test_no_exponent_notation_on_the_wire():
    body = client.get("/api/demo").text
    import re
    assert not re.search(r'"-?\d+(\.\d+)?E[+-]\d+"', body)


def test_openapi_schema_has_not_drifted():
    """web/openapi.json (source of the generated TS types) must match the live API schema (doc 07 §8 item 6)."""
    from pathlib import Path
    committed = json.loads((Path(__file__).parents[1] / "web" / "openapi.json").read_text())
    assert committed == json.loads(json.dumps(app.openapi())), \
        "Regenerate: uv run python -m rebalancer.api.export_schema && (cd web && npm run gen:types)"


def test_export_import_round_trip_keeps_data_and_strategies():
    snap = client.get("/api/demo").json()
    snap["policies"]["Klient C"]["band_abs"] = "0.4"
    x = client.post("/api/export", json={"dataset": snap["dataset"], "policies": snap["policies"]})
    assert x.status_code == 200 and x.content[:2] == b"PK"
    back = client.post("/api/import", files={"file": ("klienti.xlsx", x.content)}).json()
    assert back["policies"]["Klient C"]["band_abs"] == "0.4"
    assert back["dataset"]["clients"] == snap["dataset"]["clients"]
    c = next(c for c in back["result"]["clients"] if c["name"] == "Klient C")
    assert c["recommended"].startswith("NOW_")  # the imported strategy is applied immediately


def test_new_ticker_with_its_own_settlement_lag_round_trips():
    """A new instrument added in Excel (own currency, asset class, settlement lag) is evaluated without code changes."""
    wb = load_workbook(io.BytesIO(demo_bytes()))
    ws = wb["Instrumenty"]
    ws["G1"] = "Vypořádání (pracovní dny)"
    ws.append(["CSPX", "iShares Core S&P 500", "IE00B5BMR087", "USD", "Akcie (USA)", 610.5, 1])
    wb["Kurzy"].append(["GBP/CZK", 28.1])
    k = wb["Klienti"]
    k.insert_cols(8)
    k["H1"] = "Cíl CSPX %"
    for row in range(2, 7):
        k.cell(row, 8, 0)
    k["D3"], k["H3"] = 25, 10  # Klient B: 10 % of the 35 % equity target moves to the new ticker
    wb["Pozice"].append(["Klient B", "CSPX", datetime(2022, 5, 3), 300, 400.0])
    buf = io.BytesIO()
    wb.save(buf)
    body = client.post("/api/import", files={"file": ("novy_ticker.xlsx", buf.getvalue())}).json()
    inst = {i["ticker"]: i for i in body["dataset"]["instruments"]}
    assert inst["CSPX"]["settlement_days"] == 1 and inst["CSPX"]["currency"] == "USD"
    b = next(c for c in body["result"]["clients"] if c["name"] == "Klient B")
    assert b["status"] == "ok"
    assert any(x["key"] == "CSPX" for x in b["valuation"]["components"])
