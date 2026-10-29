"""Edge cases and adversarial inputs: validation, limits, timing and robustness of the engine, importer and API."""
from dataclasses import replace
from datetime import date
from decimal import Decimal as D
from io import BytesIO
from pathlib import Path
import os
import subprocess
import sys
import time

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import TypeAdapter

from rebalancer.adapters import demo_bytes, read_workbook, WorkbookError
from rebalancer.api.app import app
from rebalancer.engine import evaluate
from rebalancer.engine.model import CashBalance, Client, Dataset, Flow, Instrument, Lot
from rebalancer.engine.policy import default_policy, validate_policy

ASOF = date(2026, 10, 23)


def dataset(*, cash=100000, lots=(), flows=(), targets=None, currency='CZK', price='100', asof=ASOF):
    targets = targets or {'VWCE': D(50), 'EUNA': D(50), 'CASH': D(0)}
    client = Client('edge', '', 'FO', targets, lots=list(lots),
                    cash=[CashBalance('CZK', D(cash))], flows=list(flows))
    return Dataset(asof, [Instrument('VWCE', '', '', currency, 'Akcie', D(price)),
                          Instrument('EUNA', '', '', 'CZK', 'Dluhopisy', D(100))],
                   {'CZK': D(1), 'EUR': D(25), 'USD': D(22)}, [client])


def checked(ds, **policy):
    started = time.monotonic()
    result = evaluate(ds, {'edge': replace(default_policy(ds.clients[0].targets), **policy)})
    assert time.monotonic() - started < 5
    c = result.clients[0]
    assert c.status == 'ok', c.issues
    for alt in c.alternatives:
        assert all(o.units > 0 and isinstance(o.units, int) for o in alt.orders)
        bought = {o.ticker for o in alt.orders if o.side == 'BUY'}
        sold = {o.ticker for o in alt.orders if o.side == 'SELL'}
        assert not bought & sold
        assert all(ds.clients[0].targets[k] > 0 for k in bought)
        if alt.feasible:
            assert all(v >= 0 for p in alt.balances for v in p.balances.values())
        for ticker in sold:
            assert sum(o.units for o in alt.orders if o.side == 'SELL' and o.ticker == ticker) <= sum(
                lot.units for lot in ds.clients[0].lots if lot.ticker == ticker)
    return c


@pytest.mark.parametrize('inner', [D(0), D(100)])
@pytest.mark.parametrize('minimum', [D(0), D(25000)])
def test_cash_only_and_policy_edges(inner, minimum):
    c = checked(dataset(), inner_band=inner, min_trade=minimum, reserve_days=21)
    assert c.recommended is not None


@pytest.mark.parametrize('targets', [
    {'VWCE': D(0), 'EUNA': D(100), 'CASH': D(0)},
    {'VWCE': D(100), 'EUNA': D(0), 'CASH': D(0)},
    {'VWCE': D(35), 'EUNA': D(35), 'CASH': D(30)},
])
def test_zero_targets_single_instrument_and_cash_thirty(targets):
    checked(dataset(cash=50000, lots=[Lot('VWCE', date(2020, 1, 2), 1000, D(50))], targets=targets))


def test_withdrawal_larger_than_nav():
    c = checked(dataset(flows=[Flow(date(2026, 11, 2), D(-200000))]))
    assert c.outcome.reason == 'withdrawals_uncovered'
    assert c.recommended is None
    assert not any(a.feasible for a in c.alternatives)


def test_holiday_deposit_and_past_withdrawal():
    ds = dataset(flows=[Flow(date(2026, 10, 28), D(10000)), Flow(date(2026, 10, 22), D(-50000))])
    c = checked(ds)
    assert c.valuation.flows[0].past
    assert c.valuation.flows[1].value_date == date(2026, 10, 29)
    assert c.valuation.nav == D(100000)  # historical flow already reflected in starting cash
    assert any(s.reason == 'deposit' and s.date == date(2026, 10, 29)
               for a in c.alternatives if a.id == 'NOW_P' for s in a.steps)


def test_eur_only_cash_with_czk_withdrawal():
    ds = dataset(flows=[Flow(date(2026, 11, 2), D(-10000))])
    ds = replace(ds, clients=[replace(ds.clients[0], cash=[CashBalance('EUR', D(4000))])])
    c = checked(ds)
    assert any(a.feasible for a in c.alternatives)
    assert any(f.source == 'EUR' and f.target == 'CZK' for a in c.alternatives for f in a.fx)


def test_price_increase_fifty_percent(demo):
    ds, issues = demo
    ds = replace(ds, instruments=[replace(i, price=i.price * D('1.5')) if i.ticker == 'VWCE' else i
                                  for i in ds.instruments])
    r = evaluate(ds, issues=issues)
    assert all(c.status == 'ok' for c in r.clients)
    assert all(v >= 0 for c in r.clients for a in c.alternatives if a.feasible
               for p in a.balances for v in p.balances.values())


def test_exactly_funded_cash_withdrawal_is_not_uncovered():
    c = checked(dataset(targets={'VWCE': D(0), 'EUNA': D(0), 'CASH': D(100)},
                        flows=[Flow(date(2026, 11, 2), D(-100000))]))
    assert c.outcome.reason != 'withdrawals_uncovered', c.outcome


def test_future_lot_rejected_by_engine():
    ds = dataset(cash=0, lots=[Lot('VWCE', date(2027, 1, 1), 1000, D(50))])
    assert evaluate(ds).clients[0].status == 'data_error'


def test_sunday_valuation_does_not_trade_on_sunday():
    c = checked(dataset(asof=date(2026, 10, 25)))
    assert all(o.trade_date.weekday() < 5 for a in c.alternatives for o in a.orders)


def test_reserve_minimum_is_ten_business_days():
    assert validate_policy(replace(default_policy(dataset().clients[0].targets), reserve_days=10),
                           ['VWCE', 'EUNA'])


def test_valid_tiny_price_completes_in_five_seconds():
    # Separate process makes the regression bounded even when production never returns promptly.
    source = """
from datetime import date
from decimal import Decimal as D
from rebalancer.engine.model import *
from rebalancer.engine import evaluate
c=Client('x','','FO',{'VWCE':D(100),'CASH':D(0)},cash=[CashBalance('CZK',D(100000))])
ds=Dataset(date(2026,10,23),[Instrument('VWCE','','','EUR','Akcie',D('.000001'))],{'CZK':D(1),'EUR':D(25)},[c])
assert evaluate(ds).clients[0].status == 'ok'
"""
    subprocess.run([sys.executable, '-c', source], cwd=Path(__file__).resolve().parents[1],
                   env={**os.environ}, timeout=5, check=True, capture_output=True)


def workbook(change):
    wb = load_workbook(BytesIO(demo_bytes()))
    change(wb)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.parametrize('value,code', [('oops', 'bad_amount'), ('=1+2', 'formula'), (-1, 'negative_cash')])
def test_numeric_and_formula_cells_are_referenced_and_isolated(value, code):
    ds, issues = read_workbook(workbook(lambda w: w['Hotovost'].__setitem__('C6', value)))
    assert any(i.code == code and i.sheet == 'Hotovost' and i.cell == 'C6' for i in issues)
    clients = {c.name: c for c in evaluate(ds, issues=issues).clients}
    assert clients['Klient C'].status == 'data_error'
    assert clients['Klient A'].status == 'ok'


def test_missing_target_column():
    with pytest.raises(WorkbookError) as caught:
        read_workbook(workbook(lambda w: w['Klienti'].__setitem__('D1', 'missing')))
    assert any(i.code == 'missing_column' and i.sheet == 'Klienti' and i.cell for i in caught.value.issues)


@pytest.mark.parametrize('kind', ['upload', 'uncompressed', 'cells', 'lots', 'flows'])
def test_import_limits(kind):
    if kind == 'upload':
        data = b'x' * (2 * 1024 * 1024 + 1)
    elif kind == 'uncompressed':
        import zipfile
        buf = BytesIO()
        with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z:
            z.writestr('padding', b'x' * (20 * 1024 * 1024 + 1))
        data = buf.getvalue()
    elif kind == 'cells':
        data = workbook(lambda w: w['Pozice'].__setitem__('A10001', 'x'))
    else:
        def extend(w):
            ws = w['Pozice' if kind == 'lots' else 'Cash flow']
            row = [c.value for c in ws[2]]
            for _ in range(1001 if kind == 'lots' else 501):
                ws.append(row)
        data = workbook(extend)
    with pytest.raises(WorkbookError) as caught:
        read_workbook(data)
    assert caught.value.issues[0].code in {'too_large', 'too_many_cells', 'too_many_lots', 'too_many_flows'}


def test_duplicate_fx_rejected():
    with pytest.raises(WorkbookError) as caught:
        read_workbook(workbook(lambda w: w['Kurzy'].append(['EUR/CZK', 999])))
    assert any(i.code == 'duplicate' and i.sheet == 'Kurzy' and i.cell for i in caught.value.issues)


def test_duplicate_lot_rejected():
    ds, issues = read_workbook(workbook(lambda w: w['Pozice'].append([c.value for c in w['Pozice'][2]])))
    assert any(i.code == 'duplicate' and i.sheet == 'Pozice' and i.cell for i in issues)


@pytest.mark.parametrize('kind', ['orphan', 'duplicate_client'])
def test_global_import_errors_block_evaluation(kind):
    def change(w):
        if kind == 'orphan':
            w['Pozice']['A2'] = 'Unknown'
        else:
            w['Klienti'].append([c.value for c in w['Klienti'][2]])
    data = workbook(change)
    try:
        ds, issues = read_workbook(data)
    except WorkbookError:
        return
    result = evaluate(ds, issues=issues)
    assert all(c.status != 'ok' for c in result.clients), result.issues


def test_large_numeric_text_has_cell_error():
    with pytest.raises(WorkbookError) as caught:
        read_workbook(workbook(lambda w: w['Instrumenty'].__setitem__('F2', '1e100')))
    assert any(i.sheet == 'Instrumenty' and i.cell == 'F2' for i in caught.value.issues)


def payload(ds):
    return {'dataset': TypeAdapter(Dataset).dump_python(ds, mode='json'), 'policies': {}}


def test_api_decimal_strings_and_integer_units():
    response = TestClient(app).post('/api/evaluate', json=payload(dataset()))
    assert response.status_code == 200
    c = response.json()['clients'][0]
    assert isinstance(c['valuation']['nav'], str)
    assert all(isinstance(o['gross'], str) and isinstance(o['units'], int)
               for a in c['alternatives'] for o in a['orders'])


@pytest.mark.parametrize('kind', ['future_lot', 'lot_limit'])
def test_api_invalid_dataset_not_actionable(kind):
    lots = [Lot('VWCE', date(2027, 1, 1) if kind == 'future_lot' else date(2020, 1, 2), 1, D(50))]
    if kind == 'lot_limit':
        lots *= 1001
    response = TestClient(app).post('/api/evaluate', json=payload(dataset(lots=lots)))
    assert response.status_code == 422 or all(c['status'] != 'ok' for c in response.json()['clients'])


def test_chunked_api_body_limit():
    import json
    body = json.dumps({**payload(dataset()), 'padding': 'x' * (1024 * 1024)}).encode()
    response = TestClient(app).post('/api/evaluate', content=iter([body]),
                                    headers={'content-type': 'application/json'})
    assert response.status_code == 413


def test_api_bad_shape_and_declared_size_errors():
    client = TestClient(app)
    assert client.post('/api/evaluate', json={}).status_code == 422
    response = client.post('/api/evaluate', content=b'{}', headers={'content-length': str(1024 * 1024 + 1)})
    assert response.status_code == 413 and isinstance(response.json()['detail'], str)


def test_late_deposit_does_not_cover_earlier_insolvent_withdrawal():
    c = checked(dataset(targets={'VWCE': D(0), 'EUNA': D(0), 'CASH': D(100)},
                        flows=[Flow(date(2026, 11, 2), D(-200000)),
                               Flow(date(2027, 2, 1), D(300000))]))
    assert c.recommended is None
    assert c.outcome.reason == 'withdrawals_uncovered', c.outcome


def test_small_liquidity_allocations_move_to_largest():
    c = checked(dataset(cash=0, targets={'VWCE': D(35), 'EUNA': D(35), 'CASH': D(30)},
                        lots=[Lot(k, date(2020, 1, 2), 500, D(50)) for k in ['VWCE', 'EUNA']]))
    alt = next(a for a in c.alternatives if a.id == 'NOW_T')
    sales = [o for o in alt.orders if o.side == 'SELL']
    assert len(sales) == 1, [(o.ticker, o.units, o.gross_czk) for o in sales]
    assert sales[0].ticker == 'VWCE' and sales[0].units == 301
